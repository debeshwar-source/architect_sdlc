"""Plain-assert checks for the requirement document (tools/requirement_doc.py).
Run: python3 test_documents.py (or pytest).

Plain words: automatic self-checks that the Requirement agent's output turns into a
correct document, as a screen view, a Word file and a PDF - including awkward text
(symbols, markup-like text, other scripts, very long text) and the case where a
document library is not installed. Tests that need python-docx / reportlab are
skipped, with a message, where those are not installed.
"""

import sys
import zipfile
from io import BytesIO

from tools import requirement_doc as rd

try:
    import docx
    import reportlab  # noqa: F401
    HAVE_LIBS = True
except ImportError:
    HAVE_LIBS = False
try:
    import pypdf
except ImportError:
    pypdf = None


def _needs_libs(test):
    def wrapper():
        if not HAVE_LIBS:
            print("   (skipped: python-docx / reportlab not installed)")
            return
        test()
    wrapper.__name__ = test.__name__
    return wrapper


REQ = {
    "title": "Garments Website",
    "business_goal": "Sell garments online.",
    "scope_summary": "A public website to browse and order garments.",
    "acceptance_criteria": ["A visitor can list garments.", "A visitor can check out."],
    "out_of_scope": ["Mobile app", "Delivery management"],
    "build_mode": "from_scratch",
    "clarifying_questions": [{
        "question": "Take payments directly?", "why_it_matters": "Changes cost.",
        "options": ["Redirect to a provider", "Take cards in the site"],
        "recommended_option": "Redirect to a provider",
    }],
}
META = {"run_id": "run-1", "date": "13 October 2026", "knowledge_base": ""}


def _pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() for page in pypdf.PdfReader(BytesIO(data)).pages)


# ------------------------------------------------------------
# The document itself
# ------------------------------------------------------------

def test_the_document_has_the_expected_sections_in_order():
    m = rd.document_model(REQ, META)
    assert m["title"] == "Garments Website"
    assert [s["heading"] for s in m["sections"]] == ["Scope", "Acceptance criteria", "Out of scope"]
    assert dict(m["meta"])["Run"] == "run-1" and dict(m["meta"])["Type"] == "Built from scratch"
    assert dict(m["meta"])["Knowledge base"] == "none" and dict(m["meta"])["Date"] == "13 October 2026"
    criteria = m["sections"][1]["items"]
    assert criteria == [("AC-1", "A visitor can list garments."), ("AC-2", "A visitor can check out.")]
    assert "Open questions" not in rd.to_markdown(REQ, META)      # never any open questions, even if the agent sent some


def test_missing_or_empty_parts_are_handled_not_crashed_on():
    for req in ({}, None, {"title": None, "acceptance_criteria": None, "out_of_scope": [None, "", "  "],
                           "clarifying_questions": ["junk", {"question": ""}, {"options": ["x"]}]}):
        m = rd.document_model(req)
        assert m["title"] == "Requirement"
        by = {s["heading"]: s for s in m["sections"]}
        assert by["Acceptance criteria"]["items"] == [("", "None stated.")]
        assert by["Out of scope"]["items"] == ["None stated."]


def test_the_business_goal_is_not_part_of_the_document():
    # The agent still produces a business goal (the later agents use it), but it is never printed.
    goal = "UNIQUE-GOAL-TEXT sell more"
    req = {**REQ, "business_goal": goal}
    md = rd.to_markdown(req, META)
    assert goal not in md and "Business goal" not in md and "goal" not in md.lower().replace("goals", "")
    assert [s["heading"] for s in rd.document_model(req, META)["sections"]] == [
        "Scope", "Acceptance criteria", "Out of scope"]
    if HAVE_LIBS:
        d = docx.Document(BytesIO(rd.build(req, META, "docx")[0]))
        assert goal not in "\n".join(p.text for p in d.paragraphs) and "Business goal" not in [p.text for p in d.paragraphs]
        if pypdf:
            text = _pdf_text(rd.build(req, META, "pdf")[0])
            assert "UNIQUE-GOAL-TEXT" not in text and "Business goal" not in text


def test_markdown_view():
    md = rd.to_markdown(REQ, META)
    assert md.startswith("# Garments Website")
    assert "## Acceptance criteria" in md and "- **AC-1** A visitor can list garments." in md
    assert "Type: Built from scratch" in md and "Run: run-1" in md


def test_file_names():
    assert rd.suggested_filename(REQ, "pdf") == "requirement-garments-website.pdf"
    assert rd.suggested_filename({"title": "  Ünïcode / slash: ../etc  "}, "docx") == "requirement-n-code-slash-etc.docx"
    assert rd.suggested_filename({}, "pdf") == "requirement-document.pdf"
    try:
        rd.build(REQ, META, "txt")
        raise AssertionError("expected ValueError")
    except ValueError as e:
        assert "Unknown format" in str(e)


def test_control_characters_are_removed():
    m = rd.document_model({"title": "A\x00B\x07C", "scope_summary": "x\x0cy\ttab\nline"})
    assert m["title"] == "ABC"
    assert dict((s["heading"], s) for s in m["sections"])["Scope"]["text"] == "xy\ttab\nline"


# ------------------------------------------------------------
# Word
# ------------------------------------------------------------

@_needs_libs
def test_word_file_is_valid_and_has_the_content():
    data, name, mime = rd.build(REQ, META, "docx")
    assert name == "requirement-garments-website.docx" and mime.endswith("wordprocessingml.document")
    assert "word/document.xml" in zipfile.ZipFile(BytesIO(data)).namelist()
    d = docx.Document(BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "Garments Website" and "Scope" in texts and "Business goal" not in texts
    assert any(t.startswith("AC-1") and "A visitor can list garments." in t for t in texts)
    assert sum(1 for p in d.paragraphs if p.style.name == "List Bullet") == 2          # out of scope
    assert "Generated by Blueprint" in d.sections[0].footer.paragraphs[0].text
    assert d.core_properties.title == "Garments Website"


@_needs_libs
def test_word_keeps_any_script_and_survives_hostile_text():
    req = {**REQ, "title": "Résumé → 日本語 नमस्ते ₹ ✓", "scope_summary": "a\x00b <b>not bold</b> & </para> " + "word " * 5000}
    data = rd.build(req, META, "docx")[0]
    d = docx.Document(BytesIO(data))
    assert d.paragraphs[0].text == "Résumé → 日本語 नमस्ते ₹ ✓"                       # Word has no font limits
    assert any("<b>not bold</b> & </para>" in p.text and "ab" in p.text for p in d.paragraphs)


# ------------------------------------------------------------
# PDF
# ------------------------------------------------------------

@_needs_libs
def test_pdf_is_valid():
    data, name, mime = rd.build(REQ, META, "pdf")
    assert name == "requirement-garments-website.pdf" and mime == "application/pdf"
    assert data.startswith(b"%PDF") and b"%%EOF" in data[-1024:]
    if pypdf:
        reader = pypdf.PdfReader(BytesIO(data))
        assert len(reader.pages) == 1 and reader.metadata.title == "Garments Website"
        text = _pdf_text(data)
        for expected in ("Garments Website", "Scope", "AC-1", "A visitor can list garments.",
                         "Generated by Blueprint", "Page 1"):
            assert expected in text, expected
    else:
        print("   (pypdf not installed: PDF text content not checked)")


@_needs_libs
def test_pdf_shows_markup_like_text_literally_and_stands_in_for_missing_symbols():
    req = {**REQ, "title": "Tags <b>bold</b> & more",
           "scope_summary": "Use <not built here> and </para>.\nSecond line → done ✓ costs ₹5 or €3",
           "out_of_scope": ["日本語 text"]}
    data = rd.build(req, META, "pdf")[0]
    assert data.startswith(b"%PDF")                                           # nothing crashed
    if pypdf:
        text = _pdf_text(data)
        assert "Tags <b>bold</b> & more" in text and "<not built here>" in text   # shown, not interpreted
        assert "Second line -> done [x] costs Rs.5 or EUR 3" in text.replace("\n", " ") or "-> done [x]" in text
        assert "?? text" in text or "? text" in text                              # other scripts: a visible "?"


@_needs_libs
def test_a_long_requirement_runs_over_several_pdf_pages():
    req = {**REQ, "acceptance_criteria": [f"Criterion number {i}: " + "detail " * 25 for i in range(150)]}
    data = rd.build(req, META, "pdf")[0]
    if pypdf:
        reader = pypdf.PdfReader(BytesIO(data))
        assert len(reader.pages) > 3
        assert "Page 2" in _pdf_text(data) and "Criterion number 149" in _pdf_text(data)


@_needs_libs
def test_an_empty_requirement_still_makes_both_files():
    for fmt in ("docx", "pdf"):
        data = rd.build({}, {}, fmt)[0]
        assert len(data) > 1000


# ------------------------------------------------------------
# A library is not installed
# ------------------------------------------------------------

def test_a_missing_library_gives_a_clear_message():
    for fmt, module, package in (("docx", "docx", "python-docx"), ("pdf", "reportlab", "reportlab")):
        saved = {k: sys.modules.pop(k) for k in list(sys.modules) if k == module or k.startswith(module + ".")}
        sys.modules[module] = None                      # makes `import <module>` fail
        try:
            rd.build(REQ, META, fmt)
            raise AssertionError("expected MissingLibrary")
        except rd.MissingLibrary as e:
            assert f"pip install {package}" in str(e)
        finally:
            del sys.modules[module]
            sys.modules.update(saved)
    assert "Garments Website" in rd.to_markdown(REQ, META)           # the screen view needs no library


# When run directly, execute every test_ function and print "ok" for each.
if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
