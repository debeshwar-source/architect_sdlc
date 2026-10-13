"""Plain-assert checks for the Design agent's architecture, the diagrams, the design document
and publishing it to Confluence. Run: python3 test_design.py (or pytest).

Plain words: automatic self-checks that (1) the Design agent's architecture - components,
connections, flows - is complete and consistent, and a bad answer is sent back; (2) the
diagrams are drawn cleanly (no boxes on top of each other, no line through a box, nothing
off the page); (3) the design document is right in every format, and wording can be changed
without touching the design; (4) publishing to Confluence is safe (asks first, never
overwrites someone else's page, never duplicates, never leaks the token). Confluence is a
fake server; the AI is replaced by canned answers. Tests needing python-docx / reportlab /
pypdf are skipped, with a message, where those are not installed.
"""

import copy
import json
import random
import re
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from io import BytesIO
from pathlib import Path

import test_hive as t                              # temp folders, blank settings, canned-AI helpers
import config
import llm
import tools.confluence_client as cc
import tools.design_doc as dd
import tools.diagrams as dg
from agents import design as design_agent
import orchestrator.graph as graph

try:
    import docx
    import reportlab  # noqa: F401
    HAVE_LIBS = True
except ImportError:
    HAVE_LIBS = False
try:
    import pypdf
except ImportError:
    pypdf = None

PNG = b"\x89PNG\r\n\x1a\n"
KNOWN = ["api-gateway", "user-service", "notification-service", "web-app"]
TOKEN = "SECRET-CONFLUENCE-TOKEN-456"


def needs_libs(test):
    def wrapper():
        if not HAVE_LIBS:
            print("   (skipped: python-docx / reportlab not installed)")
            return
        test()
    wrapper.__name__ = test.__name__
    return wrapper


def comp(name, kind="service", status="existing", **kw):
    return {"name": name, "kind": kind, "status": status, "responsibility": f"What {name} is for.",
            "technology": "", "interfaces": [], "data_owned": [], **kw}


def conn(a, b, label="Calls", kind="sync_api", **kw):
    return {"from": a, "to": b, "kind": kind, "label": label, "protocol": "REST", "description": "", "contract": "", **kw}


def rich_design():
    """A realistic, valid design: a new service, a new database, two changed services, an
    existing gateway / client / event bus, and two outside providers."""
    return {
        "change_required": True,
        "overview": "Reminders let users ask to be told when a show is available. A new reminder-service stores them "
                    "and publishes events that notification-service turns into push and email messages.",
        "affected_applications": [
            {"name": "user-service", "role_in_change": "Expose notification preferences.", "reason": "Reminders respect them."},
            {"name": "notification-service", "role_in_change": "Consume ReminderDue events.", "reason": "It sends notifications."}],
        "new_applications_needed": [{"name": "reminder-service", "purpose": "Own reminders.", "justification": "Nothing owns scheduling."}],
        "unaffected_applications_considered": [{"name": "api-gateway", "reason": "Routing is generic."}],
        "cross_service_contracts": [{"name": "ReminderDue", "description": "Raised when a reminder is due.",
                                     "owning_application": "reminder-service", "producer": "reminder-service",
                                     "consumers": ["notification-service"]}],
        "dependency_order": ["reminder-service", "user-service", "notification-service"],
        "components": [
            comp("web-app", "client", "existing", technology="React"),
            comp("api-gateway", "gateway", "existing", technology="Kong"),
            comp("user-service", "service", "changed", interfaces=["GET /users/{id}/preferences"]),
            comp("notification-service", "service", "changed", technology="Python"),
            comp("reminder-service", "service", "new", technology="FastAPI", interfaces=["POST /reminders"], data_owned=["reminders"]),
            comp("event-bus", "messaging", "existing", technology="Kafka"),
            comp("reminders-db", "database", "new", technology="PostgreSQL"),
            comp("push-provider", "external", "external"),
            comp("email-provider", "external", "external")],
        "connections": [
            conn("web-app", "api-gateway", "Set a reminder", status="new"),
            conn("api-gateway", "reminder-service", "POST /reminders", status="new"),
            conn("reminder-service", "reminders-db", "Store reminders", "data_access", protocol="SQL", status="new"),
            conn("reminder-service", "user-service", "Read preferences", status="changed"),
            conn("reminder-service", "event-bus", "Publish ReminderDue", "async_event", protocol="Kafka topic reminders.due", status="new", contract="ReminderDue"),
            conn("event-bus", "notification-service", "Consume ReminderDue", "async_event", status="changed"),
            conn("notification-service", "push-provider", "Send push", status="existing"),
            conn("notification-service", "email-provider", "Send email", status="existing")],
        "flows": [{"name": "A user sets a reminder", "description": "The happy path.", "steps": [
            {"from": "web-app", "to": "api-gateway", "message": "POST /reminders", "kind": "call"},
            {"from": "api-gateway", "to": "reminder-service", "message": "Create reminder", "kind": "call"},
            {"from": "reminder-service", "to": "reminders-db", "message": "INSERT reminder", "kind": "call"},
            {"from": "reminder-service", "to": "api-gateway", "message": "201 Created", "kind": "return"},
            {"from": "reminder-service", "to": "event-bus", "message": "Publish ReminderDue", "kind": "async"}]}],
        "test_strategy": "Contract tests for ReminderDue and an end-to-end test.",
        "architecture_risks": ["Duplicate notifications if the event is delivered twice."],
        "knowledge_base_updates": [{"application": "reminder-service", "action": "create", "content": "# reminder-service", "reason": "New."}],
        "clarifying_questions": [{"question": "Email as well as push?", "why_it_matters": "Cost.", "options": ["Both", "Push only"], "recommended_option": "Both"}],
    }


REQ = {"title": "Show reminders", "scope_summary": "Users can ask to be reminded.", "build_mode": "extend_existing"}
META = {"run_id": "run-1", "date": "14 October 2026", "knowledge_base": "streamly"}


# ------------------------------------------------------------
# The Design agent's architecture
# ------------------------------------------------------------

def test_a_complete_architecture_passes_and_is_tidied():
    design = rich_design()
    design_agent._validate(design, KNOWN)                                  # no error
    messy = copy.deepcopy(design)
    messy["overview"] = "  A   spaced   overview \n "
    for c in messy["components"]:
        c["responsibility"] = "  " + c["responsibility"] + "  "
        for k in ("technology", "interfaces", "data_owned"):
            c.pop(k, None)                                                 # optional fields may be missing
    for k in messy["connections"]:
        for field in ("kind", "status", "protocol", "description", "contract"):
            k.pop(field, None)
    design_agent._validate(messy, KNOWN)
    design_agent._clean_architecture(messy)
    assert messy["overview"] == "A spaced overview"
    assert all(c["interfaces"] == [] and c["technology"] == "" and c["responsibility"] == c["responsibility"].strip() for c in messy["components"])
    status = {(k["from"], k["to"]): k["status"] for k in messy["connections"]}
    assert status[("web-app", "api-gateway")] == "existing"                # both ends existing
    assert status[("api-gateway", "reminder-service")] == "new"            # one end new
    assert status[("reminder-service", "user-service")] == "new"           # new beats changed
    assert status[("notification-service", "push-provider")] == "changed"  # one end changed
    assert all(k["kind"] == "sync_api" for k in messy["connections"])


def test_each_architecture_mistake_gets_an_explanation_the_ai_can_act_on():
    def broken(mutate):
        d = rich_design()
        mutate(d)
        return d

    def comp_of(d, name):
        return next(c for c in d["components"] if c["name"] == name)

    cases = [
        (lambda d: d.update(overview="  "), "overview is empty"),
        (lambda d: d.update(components=[]), "components is empty"),
        (lambda d: comp_of(d, "web-app").update(kind="router"), "has kind 'router'"),
        (lambda d: comp_of(d, "web-app").update(status="maybe"), "has status 'maybe'"),
        (lambda d: comp_of(d, "web-app").update(responsibility=""), "no responsibility"),
        (lambda d: d["components"].append(comp("web-app", "client", "existing")), "listed twice"),
        (lambda d: comp_of(d, "web-app").update(status="changed"), "marked changed but is not in affected_applications"),
        (lambda d: comp_of(d, "reminder-service").update(name="invoice-service"), "is a new application but is not in new_applications_needed"),
        (lambda d: d.update(components=[c for c in d["components"] if c["name"] != "reminder-service"]), "must appear in components with status new"),
        (lambda d: comp_of(d, "user-service").update(status="existing"), "must appear in components with status changed"),
        (lambda d: d["components"].append(comp("billing", "service", "new")), "not in new_applications_needed"),
        (lambda d: d["components"].append(comp("mystery", "service", "existing")), "not a known application"),
        (lambda d: d["connections"].append(conn("web-app", "ghost")), "not in components"),
        (lambda d: d["connections"].append(conn("web-app", "web-app")), "joins a component to itself"),
        (lambda d: d["connections"].append(conn("web-app", "user-service", label="")), "needs a label"),
        (lambda d: d["connections"].append(conn("web-app", "user-service", kind="telepathy")), "has kind 'telepathy'"),
        (lambda d: d.update(connections=[]), "connections is empty"),
        (lambda d: d["components"].append(comp("lonely", "database", "new")), "has no connection to anything"),
        (lambda d: d["flows"][0]["steps"].append({"from": "web-app", "to": "ghost", "message": "x"}), "not in components"),
        (lambda d: d["flows"][0].update(steps=[]), "has no steps"),
        (lambda d: d["flows"][0]["steps"][0].update(message=""), "step with no message"),
        (lambda d: d["flows"][0]["steps"][0].update(kind="shout"), "step of kind 'shout'"),
    ]
    for mutate, expected in cases:
        try:
            design_agent._validate(broken(mutate), KNOWN)
            raise AssertionError(f"expected an error mentioning {expected!r}")
        except ValueError as e:
            assert expected in str(e), (expected, str(e))
    # An unconnected LIBRARY is allowed (libraries are linked at build time, not called).
    ok = broken(lambda d: d["components"].append(comp("shared-lib", "library", "new")))
    design_agent._validate(ok, KNOWN)


def test_no_change_means_no_architecture():
    d = {**rich_design(), "change_required": False, "affected_applications": [], "new_applications_needed": [],
         "dependency_order": [], "overview": "Already supported."}
    design_agent._validate(d, KNOWN)                       # the architecture checks do not apply
    design_agent._clean_architecture(d)
    assert d["components"] == [] and d["connections"] == [] and d["flows"] == [] and d["overview"] == "Already supported."


def test_a_wrong_architecture_is_sent_back_with_the_reason_and_then_accepted():
    bad = rich_design()
    bad["connections"] = []                                # components but nothing joining them
    prompts = []
    with t._scripted_model([bad, rich_design()], prompts):
        result = design_agent.run({"title": "t"}, [t_doc("api-gateway"), t_doc("user-service"), t_doc("notification-service"), t_doc("web-app")])
    assert len(prompts) == 2 and "connections is empty" in prompts[1] and "YOUR PREVIOUS ANSWER WAS REJECTED" in prompts[1]
    assert len(result["connections"]) == 8 and result["connections"][0]["status"] == "new"


def t_doc(name):
    import knowledge_base.loader as kb
    return kb.ApplicationDoc(name=name, path=Path(name), content=f"# {name}", repo_url=None, purpose=None)


# ------------------------------------------------------------
# Diagrams
# ------------------------------------------------------------

def test_diagrams_are_valid_images_for_every_kind_of_input():
    d = rich_design()
    png = dg.architecture_diagram(d["components"], d["connections"])
    assert png.startswith(PNG) and len(png) > 5000
    assert dg.architecture_diagram(d["components"], d["connections"]) == png               # drawn the same every time
    for components, connections in (
        ([comp("only-one")], []),                                                           # a single box
        (d["components"], []),                                                              # boxes, no lines
        ([comp("a"), comp("b")], [conn("a", "b"), conn("b", "a", "Back"), conn("a", "b", "Again")]),   # cycles, repeats
        ([comp("a")], [conn("a", "ghost")]),                                                # lines to nothing are ignored
        ([comp("x" * 80, technology="y" * 80)], []),                                        # very long names
        ([comp("a"), comp("a")], []),                                                       # duplicate names
        ([{"name": "bare"}], []),                                                           # nothing but a name
        ([comp("price", technology="₹ / € / 日本語")], []),                                 # symbols the font lacks
        ([], []),                                                                           # nothing at all
    ):
        assert dg.architecture_diagram(components, connections).startswith(PNG), components
    wide = [comp(f"service-{i}", "service", "new") for i in range(18)]
    png = dg.architecture_diagram(wide, [conn(f"service-{i}", f"service-{(i + 5) % 18}") for i in range(18)])
    assert png.startswith(PNG)


def _overlap(a, b):
    return min(a[2], b[2]) - max(a[0], b[0]) > 3 and min(a[3], b[3]) - max(a[1], b[1]) > 3     # a visible overlap


def _random_design(seed):
    rng = random.Random(seed)
    kinds = ["service"] * 4 + ["client", "gateway", "database", "messaging", "cache", "job", "external", "library"]
    n = rng.randint(2, 14)
    comps = [{"name": f"component-{i}", "kind": rng.choice(kinds), "status": rng.choice(["new", "changed", "existing", "external"]),
              "technology": rng.choice(["", "FastAPI", "PostgreSQL"])} for i in range(n)]
    conns = []
    for _ in range(rng.randint(1, n * 2)):
        a, b = rng.sample(range(n), 2)
        conns.append({"from": f"component-{a}", "to": f"component-{b}", "kind": rng.choice(["sync_api", "async_event", "data_access"]),
                      "label": rng.choice(["Create", "Read user preferences", "Publish event", "Store", "X"]),
                      "protocol": rng.choice(["", "REST", "Kafka topic orders.created"]), "status": rng.choice(["new", "changed", "existing"])})
    return comps, conns


def test_architecture_layout_keeps_boxes_apart_and_lines_out_of_boxes():
    label_clashes = 0
    seeds = range(90)
    for seed in seeds:
        comps, conns = _random_design(seed)
        geo = {}
        dg.architecture_diagram(comps, conns, _geometry=geo)
        width, height = geo["size"]
        boxes = geo["nodes"]
        names = list(boxes)
        for i, a in enumerate(names):
            assert 0 <= boxes[a][0] and boxes[a][2] <= width and 0 <= boxes[a][1] and boxes[a][3] <= height, (seed, a, "off the page")
            for b in names[i + 1:]:
                assert not _overlap(boxes[a], boxes[b]), (seed, a, b, "boxes overlap")
        for edge, route in zip(geo["edges"], geo["routes"]):
            for name, box in boxes.items():
                if name in (edge["from"], edge["to"]):
                    continue
                assert not any(box[0] + 3 < x < box[2] - 3 and box[1] + 3 < y < box[3] - 3 for x, y in route), \
                    (seed, edge["from"], edge["to"], "line cuts through", name)
        for label in geo["labels"]:
            assert 0 <= label[0] and label[2] <= width and 0 <= label[1] and label[3] <= height, (seed, "label off the page")
            assert not any(_overlap(label, box) for box in boxes.values()), (seed, "label sits on a box")
        label_clashes += any(_overlap(a, b) for i, a in enumerate(geo["labels"]) for b in geo["labels"][i + 1:])
    assert label_clashes <= 3, f"{label_clashes} of {len(seeds)} diagrams have labels touching"   # very crowded ones only


def test_flow_and_build_order_diagrams_handle_odd_inputs():
    d = rich_design()
    flow = d["flows"][0]
    assert dg.flow_diagram(flow, d["components"]).startswith(PNG)
    assert dg.flow_diagram(flow, []).startswith(PNG)                                      # components unknown: still drawn
    assert dg.flow_diagram({"name": "empty", "steps": []}, d["components"]).startswith(PNG)
    assert dg.flow_diagram({"name": "loop", "steps": [{"from": "a", "to": "a", "message": "calls itself"}]}, []).startswith(PNG)
    long = {"name": "long", "steps": [{"from": f"s{i}", "to": f"s{i + 1}", "message": "m " * 60, "kind": "async"} for i in range(12)]}
    assert dg.flow_diagram(long, []).startswith(PNG)                                      # 13 participants, long messages
    assert dg.build_order_diagram(d["dependency_order"], d["components"]).startswith(PNG)
    assert dg.build_order_diagram([f"app-{i}" for i in range(11)], []).startswith(PNG)    # wraps over rows
    assert dg.build_order_diagram([], []).startswith(PNG)


# ------------------------------------------------------------
# The design document
# ------------------------------------------------------------

def test_the_design_document_has_the_expected_sections_and_tables():
    m = dd.build_model(REQ, rich_design(), META)
    assert m["title"] == "Show reminders — Solution Design"
    headings = [b["text"] for b in m["blocks"] if b["type"] == "h1"]
    assert headings == ["Overview", "Architecture at a glance", "Components", "Interconnections", "Key flows",
                        "Contracts between applications", "Build order", "Impact on applications", "Test strategy",
                        "Risks", "Knowledge base documentation"]
    tables = {i: b for i, b in enumerate(m["blocks"]) if b["type"] == "table"}
    components = next(b for b in tables.values() if b["header"][0] == "Component" and "Responsibility" in b["header"])
    assert len(components["rows"]) == 9 and components["rows"][4][2] == {"status": "new"}   # reminder-service
    interconnections = next(b for b in tables.values() if b["header"][:2] == ["From", "To"] and "Protocol" in b["header"])
    assert len(interconnections["rows"]) == 8
    matrix = next(b for b in tables.values() if b["header"][0] == "From \\ To")
    assert matrix["rows"][0][0] == "web-app" and "call" in matrix["rows"][0]
    assert [b["key"] for b in m["blocks"] if b["type"] == "image"] == ["architecture", "flow-1", "build-order"]
    assert dict(m["meta"])["Knowledge base"] == "streamly" and dict(m["meta"])["Type"] == "Extends an existing system"


def test_a_design_with_no_components_makes_a_small_document():
    d = {"change_required": False, "overview": "Already supported.", "components": [], "connections": [], "flows": [],
         "dependency_order": [], "affected_applications": [], "new_applications_needed": []}
    m = dd.build_model(REQ, d)
    assert [b["text"] for b in m["blocks"] if b["type"] == "h1"] == ["Overview"]
    assert dd.build_images(d) == {} and dd.attachment_names(d) == []
    assert dd.to_markdown(m).startswith("# Show reminders")


def test_editable_fields_cover_the_wording_and_nothing_structural():
    d = rich_design()
    fields = {f["path"]: f for f in dd.editable_fields(d, REQ)}
    assert fields["doc.title"]["value"] == "Show reminders — Solution Design" and fields["doc.title"]["required"]
    assert fields["overview"]["kind"] == "text" and fields["components.4.interfaces"]["value"] == ["POST /reminders"]
    assert fields["components.4.interfaces"]["kind"] == "lines" and fields["connections.0.label"]["required"]
    assert "flows.0.steps.1.message" in fields and "architecture_risks.0" in fields and "test_strategy" in fields
    # names, kinds, statuses, endpoints and order are structure: not editable
    for path in fields:
        assert not path.endswith((".name", ".kind", ".status", ".from", ".to", ".steps.0.from")) or path == "flows.0.name"
    assert "dependency_order.0" not in fields and "components.0.name" not in fields and "connections.0.from" not in fields


def test_wording_edits_change_the_document_and_diagrams_but_never_the_design():
    d, before = rich_design(), json.dumps(rich_design(), sort_keys=True)
    edits = {"doc.title": "Reminders: the design", "overview": "A rewritten overview.",
             "components.4.responsibility": "Owns reminders end to end.", "components.4.interfaces": ["POST /reminders", "GET /reminders"],
             "connections.0.label": "Create a reminder", "flows.0.steps.1.message": "Forward the request", "architecture_risks.0": "Rewritten risk."}
    m = dd.build_model(REQ, d, META, edits)
    assert m["title"] == "Reminders: the design"
    md = dd.to_markdown(m)
    for text in ("A rewritten overview.", "Owns reminders end to end.", "GET /reminders", "Create a reminder",
                 "Forward the request", "Rewritten risk."):
        assert text in md, text
    assert json.dumps(d, sort_keys=True) == before                                          # the design is untouched
    assert dd.build_images(d, REQ, edits)["architecture"] != dd.build_images(d, REQ)["architecture"]   # label edit redrawn
    assert dd.build_images(d, REQ, {"overview": "x"})["architecture"] == dd.build_images(d, REQ)["architecture"]
    # normalise_edits keeps only real changes
    kept = dd.normalise_edits(d, REQ, {"overview": d["overview"], "doc.title": "New", "components.4.interfaces": "POST /reminders"})
    assert kept == {"doc.title": "New"}


def test_wording_edits_that_are_not_allowed_are_refused():
    d = rich_design()
    for bad, expected in (({"components.0.name": "renamed"}, "cannot be edited"), ({"connections.0.from": "x"}, "cannot be edited"),
                          ({"dependency_order.0": "x"}, "cannot be edited"), ({"components.99.responsibility": "x"}, "cannot be edited"),
                          ({"overview": "  "}, "need some text"), ({"doc.title": ""}, "need some text"),
                          ({"connections.0.label": ""}, "need some text"), ({"flows.0.steps.0.message": ""}, "need some text")):
        try:
            dd.effective(d, REQ, bad)
            raise AssertionError(f"expected an error for {bad}")
        except ValueError as e:
            assert expected in str(e), (bad, str(e))
    assert dd.effective(d, REQ, {"components.0.technology": ""})[0]["components"][0]["technology"] == ""   # optional text may be cleared


def test_the_confluence_page_is_well_formed_complete_and_escaped():
    d = rich_design()
    d["overview"] = 'Uses <b>bold</b> & "quotes" in text'
    d["components"][0]["responsibility"] = "Shows <script>alert(1)</script> safely"
    pkg = dd.confluence_package(REQ, d, META, {"architecture_risks.0": "Risk & <reward>"})
    assert pkg["title"] == "Show reminders — Solution Design"
    assert sorted(pkg["attachments"]) == ["architecture.png", "build-order.png", "flow-1.png"]
    assert all(png.startswith(PNG) for png in pkg["attachments"].values())
    root = ET.fromstring('<root xmlns:ac="urn:ac" xmlns:ri="urn:ri">' + pkg["body"] + "</root>")     # well-formed
    ac, ri = "{urn:ac}", "{urn:ri}"
    attached = {e.get(ri + "filename") for e in root.iter(ri + "attachment")}
    assert attached == set(pkg["attachments"])                                       # every image is referenced, none missing
    macros = [e.get(ac + "name") for e in root.iter(ac + "structured-macro")]
    assert "toc" in macros and "info" in macros and macros.count("status") >= 9      # status lozenges in the tables
    text = "".join(root.itertext())
    assert "Uses <b>bold</b> & \"quotes\" in text" in text and "<script>" not in pkg["body"]       # shown, not run
    assert "Risk & <reward>" in text
    assert len(list(root.iter("table"))) >= 6 and any(h.text == "Components" for h in root.iter("h2"))


@needs_libs
def test_word_and_pdf_design_documents():
    d = rich_design()
    data, name, mime = dd.build(REQ, d, META, None, "docx")
    assert name == "design-show-reminders.docx" and mime.endswith("wordprocessingml.document")
    assert "word/document.xml" in zipfile.ZipFile(BytesIO(data)).namelist()
    document = docx.Document(BytesIO(data))
    assert document.paragraphs[0].text == "Show reminders — Solution Design"
    assert len(document.tables) >= 6 and len(document.inline_shapes) == 3            # tables and the three diagrams
    assert "Generated by Blueprint" in document.sections[0].footer.paragraphs[0].text
    cells = [c.text for table in document.tables for row in table.rows for c in row.cells]
    assert "NEW" in cells and "CHANGED" in cells and "reminder-service" in cells

    data, name, mime = dd.build(REQ, d, META, {"overview": "ZZ rewritten overview ZZ"}, "pdf")
    assert name == "design-show-reminders.pdf" and mime == "application/pdf"
    assert data.startswith(b"%PDF") and b"%%EOF" in data[-1024:]
    if pypdf:
        reader = pypdf.PdfReader(BytesIO(data))
        assert len(reader.pages) >= 4
        text = "\n".join(p.extract_text() for p in reader.pages)
        for expected in ("Show reminders", "ZZ rewritten overview ZZ", "Components", "Interconnections", "Key flows",
                         "reminder-service", "Publish ReminderDue", "Build order", "Page 1"):
            assert expected in text, expected
        assert sum(len(p.images) for p in reader.pages) == 3                          # the three diagrams are in the PDF
    else:
        print("   (pypdf not installed: PDF text not checked)")


def test_design_document_missing_library_message():
    for fmt, module, package in (("docx", "docx", "python-docx"), ("pdf", "reportlab", "reportlab")):
        saved = {k: sys.modules.pop(k) for k in list(sys.modules) if k == module or k.startswith(module + ".")}
        sys.modules[module] = None
        try:
            dd.build(REQ, rich_design(), META, None, fmt)
            raise AssertionError("expected MissingLibrary")
        except dd.MissingLibrary as e:
            assert f"pip install {package}" in str(e)
        finally:
            del sys.modules[module]
            sys.modules.update(saved)
    assert "Show reminders" in dd.to_markdown(dd.build_model(REQ, rich_design(), META))   # the screen view needs no library

    # reportlab is also what the DIAGRAMS use for their font: without it, Word still names the cause.
    if HAVE_LIBS:
        saved = {k: sys.modules.pop(k) for k in list(sys.modules) if k == "reportlab" or k.startswith("reportlab.")}
        sys.modules["reportlab"] = None
        import tools.fonts as fonts_module
        fonts_module._glyphs.clear()
        try:
            dd.build(REQ, rich_design(), META, None, "docx")
            raise AssertionError("expected MissingLibrary")
        except dd.MissingLibrary as e:
            assert "pip install reportlab" in str(e) and "diagrams" in str(e).lower()
        finally:
            del sys.modules["reportlab"]
            sys.modules.update(saved)
        # ...and Pillow missing says so too (rather than an import error at start-up)
        real = dg.Image
        dg.Image = None
        try:
            dg.architecture_diagram([comp("a")], [])
            raise AssertionError("expected MissingLibrary")
        except dd.MissingLibrary as e:
            assert "pip install pillow" in str(e)
        finally:
            dg.Image = real


# ------------------------------------------------------------
# Confluence, against a FAKE server
# ------------------------------------------------------------

class _Resp:
    def __init__(self, status, body=None, headers=None):
        self.status_code, self._body, self.headers = status, body if body is not None else {}, headers or {}
        self.text = json.dumps(self._body)

    def json(self):
        return self._body


class FakeConfluence:
    """Just enough of Confluence's REST API to check our client's behaviour."""

    def __init__(self, spaces=("DOCS",), existing_titles=(), fail_uploads_after=None, rate_limit_first=False,
                 bad_auth=False, forbid=False):
        self.spaces = set(spaces)
        self.pages = {}                       # id -> {"title","body","version","space","ancestors"}
        for i, title in enumerate(existing_titles, start=900):
            self.pages[str(i)] = {"title": title, "body": "someone else's", "version": 1, "space": "DOCS", "ancestors": []}
        self.attachments = {}                 # page id -> {filename: {"id","bytes","versions"}}
        self.calls, self.uploads = [], 0
        self.fail_uploads_after, self.rate_limit_first, self.bad_auth, self.forbid = fail_uploads_after, rate_limit_first, bad_auth, forbid
        self._next = 100

    def request(self, method, url, json=None, files=None, data=None, params=None, headers=None, auth=None, timeout=None):
        path = url.split("/rest/api")[1]
        self.calls.append((method, path, json, files, params, dict(headers or {}), auth))
        if self.bad_auth:
            return _Resp(401, {"message": f"bad credentials {TOKEN}"})
        if self.rate_limit_first:
            self.rate_limit_first = False
            return _Resp(429, headers={"Retry-After": "3"})
        if self.forbid:
            return _Resp(403, {"message": "no permission"})
        parts = path.strip("/").split("/")
        if method == "GET" and parts[0] == "space" and len(parts) == 2:
            return _Resp(200, {"key": parts[1]}) if parts[1] in self.spaces else _Resp(404, {"message": "No space"})
        if method == "GET" and parts == ["space"]:
            return _Resp(200, {"results": [{"key": k} for k in sorted(self.spaces)][: (params or {}).get("limit", 25)]})
        if method == "GET" and parts == ["content"]:
            hits = [{"id": i, "title": p["title"]} for i, p in self.pages.items() if p["title"] == params["title"] and p["space"] == params["spaceKey"]]
            return _Resp(200, {"results": hits[: params.get("limit", 25)]})
        if method == "POST" and parts == ["content"]:
            if any(p["title"] == json["title"] and p["space"] == json["space"]["key"] for p in self.pages.values()):
                return _Resp(400, {"message": "A page with this title already exists"})
            self._next += 1
            pid = str(self._next)
            self.pages[pid] = {"title": json["title"], "body": json["body"]["storage"]["value"], "version": 1,
                               "space": json["space"]["key"], "ancestors": [a["id"] for a in json.get("ancestors", [])]}
            return _Resp(200, {"id": pid, "_links": {"webui": f"/spaces/{json['space']['key']}/pages/{pid}/Title", "base": "https://co.atlassian.net/wiki"}})
        if parts[:1] == ["content"] and len(parts) == 2:
            page = self.pages.get(parts[1])
            if page is None:
                return _Resp(404, {"message": "no page"})
            if method == "GET":
                return _Resp(200, {"id": parts[1], "version": {"number": page["version"]}})
            if method == "PUT":
                if json["version"]["number"] != page["version"] + 1:
                    return _Resp(409, {"message": "version conflict"})
                page.update(title=json["title"], body=json["body"]["storage"]["value"], version=json["version"]["number"])
                return _Resp(200, {"id": parts[1], "_links": {"webui": f"/spaces/DOCS/pages/{parts[1]}/Title", "base": "https://co.atlassian.net/wiki"}})
        if parts[:1] == ["content"] and parts[2:3] == ["child"] and parts[3:4] == ["attachment"]:
            store = self.attachments.setdefault(parts[1], {})
            if method == "GET":
                name = (params or {}).get("filename")
                return _Resp(200, {"results": [{"id": a["id"], "title": n} for n, a in store.items() if n == name]})
            if method == "POST":
                if headers.get("X-Atlassian-Token") != "no-check":
                    return _Resp(403, {"message": "XSRF check failed"})
                if self.fail_uploads_after is not None and self.uploads >= self.fail_uploads_after:
                    return _Resp(500 + 3 - 3, {"message": "boom"}) if False else _Resp(403, {"message": "cannot attach"})
                self.uploads += 1
                filename, content, _ = files["file"]
                if len(parts) == 4:
                    store[filename] = {"id": f"att{len(store) + 1}", "bytes": content, "versions": 1}
                else:
                    store[filename]["bytes"] = content
                    store[filename]["versions"] += 1
                return _Resp(200, {"results": [{"id": store[filename]["id"]}]})
        return _Resp(404, {"message": "unknown"})


def client(fake, email="me@co.com", parent=""):
    waits = []
    c = cc.ConfluenceClient("https://co.atlassian.net", "DOCS", TOKEN, email, parent_id=parent, session=fake, sleep=waits.append)
    c.waits = waits
    return c


PAGE = {"title": "Show reminders — Solution Design", "body": "<p>Body</p>", "files": {"architecture.png": PNG + b"1", "flow-1.png": PNG + b"2"}}


def publish(c, state_path=None, **over):
    p = {**PAGE, **over}
    return c.publish(p["title"], p["body"], p["files"], "run-1", state_path)


def confluence_env(on=True):
    import os
    names = ("JIRA_BASE_URL", "JIRA_EMAIL", "JIRA_API_TOKEN", "CONFLUENCE_SPACE_KEY", "CONFLUENCE_PARENT_PAGE_ID",
             "CONFLUENCE_BASE_URL", "CONFLUENCE_EMAIL", "CONFLUENCE_API_TOKEN")
    saved = {n: os.environ.get(n) for n in names}
    values = {"JIRA_BASE_URL": "https://co.atlassian.net", "JIRA_EMAIL": "me@co.com", "JIRA_API_TOKEN": TOKEN,
              "CONFLUENCE_SPACE_KEY": "DOCS"} if on else {}
    os.environ.update({n: values.get(n, "") for n in names})

    class _Ctx:
        def __enter__(self): return self
        def __exit__(self, *exc):
            for n, v in saved.items():
                os.environ[n] = v or ""
    return _Ctx()


def test_confluence_reuses_the_jira_credentials_and_the_preview_is_offline_and_secret_free():
    with confluence_env(False):
        p = cc.preview("T", 3)
        assert not p["configured"] and "CONFLUENCE_SPACE_KEY" in " ".join(p["missing"])
    with confluence_env():
        p = cc.preview("My page", 3)
        assert p["configured"] and p["base_url"] == "https://co.atlassian.net/wiki"      # /wiki added for Cloud
        assert p["space_key"] == "DOCS" and p["account"] == "me@co.com" and p["attachments"] == 3
        assert TOKEN not in repr(p)
        import os
        os.environ["CONFLUENCE_BASE_URL"] = "https://wiki.example.com/confluence"          # Data Center: used as given
        os.environ["CONFLUENCE_EMAIL"] = ""
        assert cc.preview()["base_url"] == "https://wiki.example.com/confluence"
        os.environ["CONFLUENCE_SPACE_KEY"] = ""
        assert not cc.is_configured()
    with confluence_env():
        assert cc.from_env(session=FakeConfluence())._email == "me@co.com"


def test_publish_creates_the_page_under_its_parent_and_attaches_every_image():
    fake = FakeConfluence()
    result = publish(client(fake, parent="555"))
    page = fake.pages[result["page_id"]]
    assert page["title"] == PAGE["title"] and page["body"] == PAGE["body"] and page["ancestors"] == ["555"] and page["space"] == "DOCS"
    assert sorted(fake.attachments[result["page_id"]]) == ["architecture.png", "flow-1.png"]
    assert result["status"] == "published" and result["url"] == f"https://co.atlassian.net/wiki/spaces/DOCS/pages/{result['page_id']}/Title"
    assert not result["updated"] and result["attachments"] == ["architecture.png", "flow-1.png"]
    uploads = [c for c in fake.calls if c[0] == "POST" and "attachment" in c[1]]
    assert uploads and all(c[5].get("X-Atlassian-Token") == "no-check" for c in uploads)   # required by Confluence
    assert fake.calls[0][6] == ("me@co.com", TOKEN) and "Authorization" not in fake.calls[0][5]   # Cloud: email + token


def test_data_center_uses_a_bearer_token():
    fake = FakeConfluence()
    publish(client(fake, email=""))
    assert fake.calls[0][6] is None and fake.calls[0][5]["Authorization"] == f"Bearer {TOKEN}"


def test_publishing_again_does_nothing_and_after_edits_updates_the_same_page():
    fake, state = FakeConfluence(), Path(tempfile.mkdtemp()) / "confluence_created.json"
    first = publish(client(fake), state)
    n = len(fake.calls)
    again = publish(client(fake), state)                                              # same document
    assert len(fake.calls) == n and again["unchanged"] and again["page_id"] == first["page_id"]

    edited = publish(client(fake), state, title="Reworded title", body="<p>New body</p>", files={"architecture.png": PNG + b"changed", "flow-1.png": PNG + b"2"})
    assert edited["updated"] and edited["page_id"] == first["page_id"] and len(fake.pages) == 1      # no second page
    page = fake.pages[first["page_id"]]
    assert page["title"] == "Reworded title" and page["body"] == "<p>New body</p>" and page["version"] == 2
    assert fake.attachments[first["page_id"]]["architecture.png"]["bytes"].endswith(b"changed")
    assert fake.attachments[first["page_id"]]["architecture.png"]["versions"] == 2           # replaced, not duplicated
    assert len(fake.attachments[first["page_id"]]) == 2


def test_publishing_never_takes_over_a_page_someone_else_wrote():
    fake = FakeConfluence(existing_titles=[PAGE["title"]])
    try:
        publish(client(fake))
        raise AssertionError("expected ConfluenceError")
    except cc.ConfluenceError as e:
        assert "already exists" in str(e) and "different title" in str(e)
    assert len(fake.pages) == 1 and fake.pages["900"]["body"] == "someone else's"             # untouched
    assert not any(c[0] == "POST" and c[1] == "/content" for c in fake.calls)                  # nothing was created


def test_publishing_resumes_after_an_upload_fails_without_a_second_page():
    fake, state = FakeConfluence(fail_uploads_after=1), Path(tempfile.mkdtemp()) / "confluence_created.json"
    try:
        publish(client(fake), state)
        raise AssertionError("expected ConfluenceError")
    except cc.ConfluenceError as e:
        assert "permission" in str(e) and TOKEN not in str(e)
    assert len(fake.pages) == 1 and len(fake.attachments["101"]) == 1 and state.is_file()     # the page and one image exist
    fake.fail_uploads_after = None
    result = publish(client(fake), state)
    assert len(fake.pages) == 1 and sorted(fake.attachments["101"]) == ["architecture.png", "flow-1.png"]
    assert fake.attachments["101"]["architecture.png"]["versions"] == 1                          # not re-sent
    assert result["status"] == "published"


def test_confluence_problems_are_explained_in_plain_words_and_never_leak_the_token():
    def failure(fake):
        try:
            publish(client(fake))
        except cc.ConfluenceError as e:
            return str(e)
        raise AssertionError("expected ConfluenceError")

    msg = failure(FakeConfluence(spaces=("OTHER",)))
    assert "Space 'DOCS' was not found" in msg and "CONFLUENCE_SPACE_KEY" in msg
    msg = failure(FakeConfluence(spaces=()))
    assert "cannot see ANY space" in msg
    msg = failure(FakeConfluence(bad_auth=True))
    assert "credentials were rejected" in msg and TOKEN not in msg
    msg = failure(FakeConfluence(forbid=True))
    assert "lacks permission" in msg and "DOCS" in msg
    try:
        client(FakeConfluence()).publish("   ", "<p/>", {}, "r")
        raise AssertionError("expected ConfluenceError")
    except cc.ConfluenceError as e:
        assert "needs a title" in str(e)


def test_confluence_waits_out_a_rate_limit_and_retries():
    fake = FakeConfluence(rate_limit_first=True)
    c = client(fake)
    publish(c)
    assert c.waits and c.waits[0] == 3.0 and len(fake.pages) == 1


# ------------------------------------------------------------
# The pipeline step
# ------------------------------------------------------------

def _state(**over):
    d = rich_design()
    base = {"run_id": graph.new_run_id(), "requirement": REQ, "design": d, "iteration": 1, "history": []}
    base.update(over)
    return base


def _with_fake_confluence(fake):
    original = cc.from_env
    cc.from_env = lambda **kw: client(fake)

    class _Ctx:
        def __enter__(self): return fake
        def __exit__(self, *exc): setattr(cc, "from_env", original)
    return _Ctx()


def test_the_confluence_step_asks_first_and_publishes_only_when_approved():
    # not configured -> skipped; nothing built or sent
    r = graph.confluence_publish_node(_state())["confluence_result"]
    assert r["status"] == "skipped" and "CONFLUENCE_SPACE_KEY" in r["reason"]

    with confluence_env(), _with_fake_confluence(FakeConfluence()) as fake:
        for state in (_state(), _state(confluence_approved=False), _state(confluence_confirm=lambda p: False)):
            out = graph.confluence_publish_node(state)
            assert out["confluence_result"]["status"] == "declined" and fake.calls == []
        asked = []
        out = graph.confluence_publish_node(_state(confluence_confirm=lambda p: asked.append(p) or True))
        assert out["confluence_result"]["status"] == "published" and len(fake.pages) == 1
        assert asked[0]["title"] == "Show reminders — Solution Design" and asked[0]["attachments"] == 3 and TOKEN not in repr(asked[0])
        page = next(iter(fake.pages.values()))
        assert "Reminders let users ask" in page["body"] and sorted(next(iter(fake.attachments.values()))) == ["architecture.png", "build-order.png", "flow-1.png"]


def test_the_confluence_step_publishes_the_users_wording_and_records_failures():
    edits = {"doc.title": "Reminders design (reworded)", "overview": "A reworded overview."}
    with confluence_env(), _with_fake_confluence(FakeConfluence()) as fake:
        out = graph.confluence_publish_node(_state(design_doc_edits=edits, confluence_approved=True))
        page = next(iter(fake.pages.values()))
        assert page["title"] == "Reminders design (reworded)" and "A reworded overview." in page["body"]
        assert out["confluence_result"]["title"] == "Reminders design (reworded)"
    with confluence_env(), _with_fake_confluence(FakeConfluence(spaces=("OTHER",))):
        out = graph.confluence_publish_node(_state(confluence_approved=True))              # a failure does not crash
        assert out["confluence_result"]["status"] == "failed" and "not found" in out["confluence_result"]["error"]


def test_there_is_nothing_to_publish_when_no_change_was_needed():
    state = _state(design={"change_required": False, "overview": "Already supported.", "components": []})
    with confluence_env(), _with_fake_confluence(FakeConfluence()) as fake:
        r = graph.confluence_publish_node({**state, "confluence_approved": True})["confluence_result"]
    assert r["status"] == "skipped" and fake.calls == []


def test_design_doc_edits_are_validated_and_a_new_design_starts_clean():
    state = _state()
    update = graph.apply_design_doc_edits(state, {"overview": "Reworded.", "doc.title": state["design"]["overview"][:0] or "New title"})
    assert update["design_doc_edits"] == {"overview": "Reworded.", "doc.title": "New title"}
    assert update["history"][-1]["result"]["changes"] == ["doc.title", "overview"] and update["confluence_approved"] is None
    assert graph.apply_design_doc_edits({**state, **update}, {"overview": "Reworded.", "doc.title": "New title"}) == {}
    assert graph.apply_design_doc_edits(state, {"overview": state["design"]["overview"]}) == {}      # typed back = no change
    for bad in ({"components.0.name": "x"}, {"overview": ""}):
        try:
            graph.apply_design_doc_edits(state, bad)
            raise AssertionError("expected ValueError")
        except ValueError:
            pass

    # a re-run design resets the edits (they pointed at parts of the old design)
    canned = t._scratch_canned("from_scratch")
    canned["Design Agent"] = t._design()
    log = []
    original = llm.call_llm_json
    llm.call_llm_json = t._fake_llm(canned, log)
    try:
        out = graph.design_node({"run_id": graph.new_run_id(), "requirement_text": "t", "requirement": {"title": "t"},
                                 "applications": [], "iteration": 0, "history": [], "design_doc_edits": {"overview": "old"},
                                 "confluence_approved": True, "confluence_result": {"status": "published"}})
    finally:
        llm.call_llm_json = original
    assert out["design_doc_edits"] == {} and out["confluence_approved"] is None and out["confluence_result"] == {}


def test_the_whole_pipeline_ends_with_the_confluence_step():
    canned = t._scratch_canned("from_scratch")
    canned["Design Agent"] = t._design()
    with confluence_env(), _with_fake_confluence(FakeConfluence()) as fake:
        final, log = t._run_pipeline([], "Build a todo app", canned, force_from_scratch=True,
                                     confluence_confirm=lambda p: True)
    assert final["confluence_result"]["status"] == "published" and len(fake.pages) == 1
    assert final["run_outcome"] == "APPROVED_RELEASED"                                         # the outcome is unchanged
    assert [h["stage"] for h in final["history"]][-3:] == ["kb_save", "confluence_publish", "finalize"]


# When run directly, execute every test_ function and print "ok" for each.
if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
