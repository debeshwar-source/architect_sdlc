"""
The Quality Planner agent's test plan as an Excel workbook (.xlsx).

Plain words: the agent produces structured test cases. This lays them out as a
workbook a tester can open and use straight away:

    Summary      what the plan is for, and how many cases of each type / priority
    Test Cases   one row per case: steps, test data, expected result, and a Status
                 column (with a drop-down) for the tester to fill in
    Coverage     every acceptance criterion and the test cases that prove it

Nothing here needs an account, and nothing is sent anywhere: the file is made on
this computer. openpyxl is imported only when the file is made, so the rest of the
app still works without it; `build()` then raises MissingLibrary with the fix.
"""

import math
import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from tools.fonts import MissingLibrary

FORMATS = {
    "xlsx": ("Excel workbook (.xlsx)", "xlsx",
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}

STATUSES = ["Not run", "Pass", "Fail", "Blocked"]
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")      # characters a workbook cannot hold

# (header, key, column width in characters)
COLUMNS = [
    ("ID", "id", 9), ("Title", "title", 34), ("Component", "component", 18), ("Type", "type", 13),
    ("Priority", "priority", 10), ("Covers", "refs", 11), ("Preconditions", "preconditions", 30),
    ("Steps", "steps", 46), ("Test data", "test_data", 24), ("Expected result", "expected_result", 40),
    ("Status", "status", 11), ("Tester notes", "notes", 28),
]


def _columns(rows: List[Dict[str, str]]) -> List[Tuple[str, str, int]]:
    # Preconditions and test data are optional: a column nobody filled in is left out.
    optional = {"preconditions", "test_data"}
    return [c for c in COLUMNS if c[1] not in optional or any(r[c[1]] for r in rows)]


def _clean(value: Any) -> str:
    return _CONTROL.sub("", "" if value is None else str(value)).strip()


def _rows(plan: Dict[str, Any]) -> List[Dict[str, str]]:
    rows = []
    for case in plan.get("test_cases") or []:
        if not isinstance(case, dict):
            continue
        steps = [_clean(s) for s in case.get("steps") or [] if _clean(s)]
        rows.append({
            "id": _clean(case.get("id")), "title": _clean(case.get("title")),
            "component": _clean(case.get("component")), "type": _clean(case.get("type")),
            "priority": _clean(case.get("priority")),
            "refs": ", ".join(_clean(r) for r in case.get("requirement_refs") or [] if _clean(r)),
            "preconditions": _clean(case.get("preconditions")),
            "steps": "\n".join(f"{i}. {s}" for i, s in enumerate(steps, start=1)),
            "test_data": _clean(case.get("test_data")),
            "expected_result": _clean(case.get("expected_result")),
            "status": STATUSES[0], "notes": "",
        })
    return rows


def _criteria(requirement: Dict[str, Any]) -> List[Tuple[str, str]]:
    items = [_clean(c) for c in (requirement or {}).get("acceptance_criteria") or [] if _clean(c)]
    return [(f"AC-{i}", text) for i, text in enumerate(items, start=1)]


def suggested_filename(requirement: Dict[str, Any]) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", _clean((requirement or {}).get("title")).lower()).strip("-")[:50]
    return f"test-plan-{slug or 'document'}.xlsx"


def title_of(requirement: Dict[str, Any]) -> str:
    return f"{_clean((requirement or {}).get('title')) or 'Requirement'} — Test Plan"


def counts(plan: Dict[str, Any]) -> Dict[str, Dict[str, int]]:
    """How many cases of each type and priority (in a sensible order, zeros left out)."""
    from agents.quality import CASE_TYPES, PRIORITIES
    rows = _rows(plan)
    by = lambda key, order: {name: n for name in order if (n := sum(1 for r in rows if r[key] == name))}
    return {"type": by("type", CASE_TYPES), "priority": by("priority", PRIORITIES)}


def coverage(requirement: Dict[str, Any], plan: Dict[str, Any]) -> List[Dict[str, Any]]:
    """For each acceptance criterion: the ids of the test cases that cover it."""
    rows = _rows(plan)
    return [{"label": label, "criterion": text,
             "cases": [r["id"] for r in rows if label in [x.strip() for x in r["refs"].split(",")]]}
            for label, text in _criteria(requirement)]


def to_xlsx(requirement: Dict[str, Any], plan: Dict[str, Any], meta: Optional[Dict[str, Any]] = None) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise MissingLibrary("Excel export needs openpyxl. Install it with: pip install openpyxl") from exc

    meta = meta or {}
    rows = _rows(plan or {})
    title = title_of(requirement)

    font = "Calibri"
    ink, grey = "1F2937", "6B7280"
    head_fill = PatternFill("solid", fgColor="1F2937")
    band_fill = PatternFill("solid", fgColor="F3F4F6")
    line = Side(style="thin", color="D1D5DB")
    border = Border(left=line, right=line, top=line, bottom=line)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    f_body = Font(name=font, size=10, color=ink)
    f_head = Font(name=font, size=10, bold=True, color="FFFFFF")
    f_title = Font(name=font, size=16, bold=True, color=ink)
    f_label = Font(name=font, size=10, bold=True, color=grey)
    priority_color = {"Critical": "B91C1C", "High": "C2410C", "Medium": "1F2937", "Low": grey}

    wb = Workbook()
    wb.properties.title = title
    wb.properties.creator = "Blueprint"

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 78
    ws["A1"] = title
    ws["A1"].font = f_title
    ws.row_dimensions[1].height = 26
    info = [("Date", _clean(meta.get("date")) or datetime.now(timezone.utc).strftime("%d %B %Y"))]
    if meta.get("run_id"):
        info.append(("Run", _clean(meta["run_id"])))
    info += [("Test cases", str(len(rows))), ("Acceptance criteria", str(len(_criteria(requirement))))]
    r = 3
    for label, value in info:
        ws.cell(r, 1, label).font = f_label
        ws.cell(r, 2, value).font = f_body
        ws.cell(r, 2).alignment = Alignment(horizontal="left")
        r += 1
    r += 1
    ws.cell(r, 1, "Approach").font = f_label
    ws.cell(r, 2, _clean((plan or {}).get("summary")) or "Not stated.").font = f_body
    ws.cell(r, 2).alignment = wrap_top
    ws.row_dimensions[r].height = max(15, 14 * math.ceil(len(_clean((plan or {}).get("summary"))) / 85))
    r += 2
    for heading, key in (("By type", "type"), ("By priority", "priority")):
        for col, text in ((1, heading), (2, "Cases")):
            c = ws.cell(r, col, text)
            c.font, c.fill, c.border = f_head, head_fill, border
        r += 1
        for name, n in counts(plan or {})[key].items():
            ws.cell(r, 1, name).font = f_body
            ws.cell(r, 2, n).font = f_body
            ws.cell(r, 2).alignment = Alignment(horizontal="left")
            for col in (1, 2):
                ws.cell(r, col).border = border
            r += 1
        r += 1

    # ---------------- Test Cases ----------------
    tc = wb.create_sheet("Test Cases")
    columns = _columns(rows)
    for col, (header, _, width) in enumerate(columns, start=1):
        c = tc.cell(1, col, header)
        c.font, c.fill, c.border = f_head, head_fill, border
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        tc.column_dimensions[get_column_letter(col)].width = width
    tc.row_dimensions[1].height = 24
    for n, row in enumerate(rows, start=2):
        lines = 1
        for col, (_, key, width) in enumerate(columns, start=1):
            value = row[key]
            c = tc.cell(n, col, value)
            c.font = Font(name=font, size=10, color=priority_color.get(value, ink), bold=(key == "priority" and value == "Critical")) \
                if key == "priority" else f_body
            c.alignment, c.border = wrap_top, border
            if n % 2 == 1:
                c.fill = band_fill
            lines = max(lines, sum(max(1, math.ceil(len(part) / max(width - 2, 1))) for part in value.split("\n")))
        tc.row_dimensions[n].height = min(max(18, 13.5 * lines + 4), 400)
    tc.freeze_panes = "C2"
    if rows:
        tc.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        status_col = get_column_letter([k for _, k, _ in columns].index("status") + 1)
        check = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        tc.add_data_validation(check)
        check.add(f"{status_col}2:{status_col}{len(rows) + 1}")
    tc.page_setup.orientation = "landscape"
    tc.page_setup.fitToWidth = 1
    tc.page_setup.fitToHeight = 0
    tc.sheet_properties.pageSetUpPr.fitToPage = True
    tc.print_title_rows = "1:1"

    # ---------------- Coverage ----------------
    cv = wb.create_sheet("Coverage")
    for col, (header, width) in enumerate((("Criterion", 11), ("Acceptance criterion", 70), ("Test cases", 40), ("Count", 8)), start=1):
        c = cv.cell(1, col, header)
        c.font, c.fill, c.border = f_head, head_fill, border
        cv.column_dimensions[get_column_letter(col)].width = width
    for n, item in enumerate(coverage(requirement, plan or {}), start=2):
        values = (item["label"], item["criterion"], ", ".join(item["cases"]) or "NOT COVERED", len(item["cases"]))
        for col, value in enumerate(values, start=1):
            c = cv.cell(n, col, value)
            c.font = Font(name=font, size=10, color="B91C1C", bold=True) if (col == 3 and not item["cases"]) else f_body
            c.alignment, c.border = wrap_top, border
        cv.row_dimensions[n].height = max(18, 13.5 * math.ceil(len(item["criterion"]) / 68) + 4)
    cv.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build(requirement: Dict[str, Any], plan: Dict[str, Any], meta: Optional[Dict[str, Any]] = None,
          fmt: str = "xlsx") -> Tuple[bytes, str, str]:
    """Make the workbook. Returns (file bytes, suggested file name, MIME type). Raises
    MissingLibrary if openpyxl is absent."""
    if fmt not in FORMATS:
        raise ValueError(f"Unknown format {fmt!r}; choose one of {sorted(FORMATS)}")
    return to_xlsx(requirement, plan, meta), suggested_filename(requirement), FORMATS[fmt][2]
